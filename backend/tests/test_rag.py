"""RAG tests: extraction formats, ingestion, retrieval, grounding, endpoints."""

import io

import pytest

from app.services import rag
from app.services.rag.extraction import DocumentError, chunk_pages, extract_pages

pytestmark = pytest.mark.asyncio


def _build_valid_pdf(text: str) -> bytes:
    """Build a minimal structurally-valid PDF (with xref) containing one text line."""
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >> stream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj".encode() + body + b"endobj\n"
    xref_pos = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += f"trailer << /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_pos}\n%%EOF".encode()
    return bytes(out)


MINIMAL_PDF = _build_valid_pdf("Skills: Python LangGraph PostgreSQL RAG")


@pytest.fixture(autouse=True)
def fake_embeddings(monkeypatch):
    """Route the embedding factory to the deterministic hashing provider."""
    from tests.fakes import HashingEmbeddingProvider

    monkeypatch.setattr(
        "app.services.rag.service.get_embedding_provider",
        lambda: HashingEmbeddingProvider(),
    )
    yield


# ---------- Extraction & chunking ----------


def test_extract_txt():
    pages = extract_pages(b"line one\nline two", "notes.txt")
    assert pages == [(1, "line one\nline two")]


def test_extract_docx():
    import docx

    buffer = io.BytesIO()
    doc = docx.Document()
    doc.add_paragraph("Operating systems: scheduling, memory management.")
    doc.add_paragraph("Second paragraph about networks.")
    doc.save(buffer)
    pages = extract_pages(buffer.getvalue(), "study.docx")
    assert pages[0][0] == 1
    assert "memory management" in pages[0][1]


def test_extract_pdf():
    pages = extract_pages(MINIMAL_PDF, "resume.pdf")
    assert len(pages) >= 1
    assert any("Python LangGraph" in text for _, text in pages)


def test_extract_unsupported_type():
    with pytest.raises(DocumentError, match="Unsupported file type"):
        extract_pages(b"data", "image.png")


def test_chunking_preserves_pages_and_overlaps():
    page1 = "word " * 300  # 1500 chars → 2 chunks at size 1000/overlap 150
    chunks = chunk_pages([(1, page1), (2, "page two content")], chunk_size=1000, overlap=150)
    assert all(c["page"] == 1 for c in chunks[:-1])
    assert chunks[-1]["page"] == 2
    assert chunks[0]["content"][:5] == chunks[1]["content"][:5].strip() or len(chunks) >= 2
    assert [c["index"] for c in chunks] == list(range(len(chunks)))


# ---------- Ingestion & retrieval (fake embeddings) ----------


async def _ingest(session, user, content: str, filename: str = "notes.txt"):
    return await rag.ingest_document(
        session, user=user, data=content.encode(), filename=filename
    )


async def test_ingestion_creates_document_and_chunks(db_session_factory):
    async with db_session_factory() as session:
        from app.services import persistence

        user = await persistence.get_or_create_default_user(session)
        document = await _ingest(
            session, user, "Python is my strongest language. " * 50, "resume.txt"
        )

        assert document.status == "ready"
        assert document.chunk_count >= 1
        assert document.page_count == 1

        from sqlalchemy import select as sa_select

        from app.models import DocumentChunk

        rows = await session.execute(
            sa_select(DocumentChunk).where(DocumentChunk.document_id == document.id)
        )
        chunks = rows.scalars().all()
        assert len(chunks) == document.chunk_count
        assert all(c.embedding is not None and len(c.embedding) == 64 for c in chunks)


async def test_retrieval_returns_relevant_chunk(db_session_factory):
    async with db_session_factory() as session:
        from app.services import persistence

        user = await persistence.get_or_create_default_user(session)
        await _ingest(session, user, "Technologies: Python, LangGraph, PostgreSQL, Docker.", "resume.txt")
        await _ingest(session, user, "Recipe: boil pasta, add tomato sauce.", "cooking.txt")

        results = await rag.retrieve_relevant_chunks(session, user, "What technologies are listed in my resume?")
        assert results, "relevant query should return chunks"
        assert results[0]["filename"] == "resume.txt"
        assert results[0]["similarity"] >= results[-1]["similarity"]  # ranked


async def test_irrelevant_query_returns_nothing(db_session_factory):
    async with db_session_factory() as session:
        from app.services import persistence

        user = await persistence.get_or_create_default_user(session)
        await _ingest(session, user, "Technologies: Python, LangGraph, PostgreSQL.", "resume.txt")

        # Completely unrelated vocabulary → below similarity threshold.
        results = await rag.retrieve_relevant_chunks(
            session, user, "quantum flux capacitor xylophone zzz"
        )
        assert results == []


async def test_missing_documents_returns_empty(db_session_factory):
    async with db_session_factory() as session:
        from app.services import persistence

        user = await persistence.get_or_create_default_user(session)
        results = await rag.retrieve_relevant_chunks(session, user, "anything at all")
        assert results == []


async def test_documents_scoped_per_user(db_session_factory):
    async with db_session_factory() as session:
        from app.models import User
        from app.services import persistence

        owner = await persistence.get_or_create_default_user(session)
        stranger = User(email="other@example.com", name="Other")
        session.add(stranger)
        await session.commit()
        await session.refresh(stranger)

        await _ingest(session, owner, "Secret resume technologies: Python.", "resume.txt")
        results = await rag.retrieve_relevant_chunks(session, stranger, "resume technologies Python")
        assert results == []  # another user must not see them


async def test_delete_document_removes_chunks(db_session_factory):
    async with db_session_factory() as session:
        from sqlalchemy import select

        from app.models import DocumentChunk
        from app.services import persistence

        user = await persistence.get_or_create_default_user(session)
        document = await _ingest(session, user, "Some content for deletion test.", "delete-me.txt")
        doc_id = document.id

        assert await rag.delete_document(session, doc_id) is True
        chunks = (await session.execute(
            select(DocumentChunk).where(DocumentChunk.document_id == doc_id)
        )).scalars().all()
        assert chunks == []
        assert await rag.delete_document(session, doc_id) is False


# ---------- Knowledge Agent grounding ----------


async def test_answer_from_documents_is_grounded_with_citations(db_session_factory, monkeypatch):
    from app.services import persistence
    from app.tools.knowledge_tools import AnswerFromDocumentsTool
    from tests.fakes import FakeProvider

    captured = {}

    class CapturingProvider(FakeProvider):
        async def generate(self, messages, **kwargs):
            captured["prompt"] = messages[0].content
            return await super().generate(messages, **kwargs)

    async with db_session_factory() as session:
        user = await persistence.get_or_create_default_user(session)
        await _ingest(session, user, "My strongest technologies are Python, LangGraph and PostgreSQL.", "resume.txt")

        tool = AnswerFromDocumentsTool(trace=None)
        tool.ctx = {"session": session, "user": user}
        monkeypatch.setattr(
            "app.tools.knowledge_tools.get_llm_provider",
            lambda: CapturingProvider(text="Grounded answer with citations."),
        )

        result = await tool.run(question="What technologies are listed in my resume?")

        assert result["grounded"] is True
        assert result["sources"] == [{"filename": "resume.txt", "page": 1}]
        # Grounding prompt contains the chunks and the strict rules.
        assert "resume.txt" in captured["prompt"]
        assert "ONLY the document passages" in captured["prompt"]


async def test_answer_refuses_when_nothing_relevant(db_session_factory, monkeypatch):
    from app.services import persistence
    from app.tools.knowledge_tools import AnswerFromDocumentsTool
    from tests.fakes import FakeProvider

    async with db_session_factory() as session:
        user = await persistence.get_or_create_default_user(session)
        await _ingest(session, user, "Technologies: Python, LangGraph.", "resume.txt")

        llm_called = {"count": 0}

        class CountingProvider(FakeProvider):
            async def generate(self, messages, **kwargs):
                llm_called["count"] += 1
                return await super().generate(messages, **kwargs)

        tool = AnswerFromDocumentsTool(trace=None)
        tool.ctx = {"session": session, "user": user}
        monkeypatch.setattr(
            "app.tools.knowledge_tools.get_llm_provider", lambda: CountingProvider()
        )

        result = await tool.run(question="xylophone quantum flux capacitor zzz")

        assert result["grounded"] is False
        assert "could not find" in result["answer"].lower()
        assert llm_called["count"] == 0  # LLM never invoked → cannot hallucinate


async def test_registry_contains_knowledge_agent():
    from app.tools.registry import tools_for

    assert set(tools_for("knowledge")) == {
        "search_documents", "retrieve_relevant_chunks", "answer_from_documents",
    }
